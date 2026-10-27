"""H444 (review round 2) — a DOCX read says what it holds, and partial reads are honest.

* A DOCX is read in document order — paragraphs *and* tables (as ``|`` rows),
  text boxes, headers/footers and foot/endnotes — so content kept in a table is not
  silently dropped into an apparently complete 5-byte read.
* Embedded images, charts, SmartArt and OLE objects are named as ``visual_content``
  (no OCR); equations, which are not extracted, are named as a ``coverage_warning``.
* An extraction stopped by the wall-clock budget is not cached: the next read parses
  again and can return the whole document. A text-budget stop (deterministic) is.
* A kanban worker's per-call file handler shares the owner's document cache.
* ``offset: 0`` beside a line window is the absent offset, not a conflict, and a line
  page never ends inside a UTF-8 character.
"""

import importlib.machinery
import io
import sys
import xml.etree.ElementTree as ET
import zlib
from types import ModuleType, SimpleNamespace

import pytest

from agents.core import file_tools, local_docs
from agents.core.file_tools import FileScope, FileTools, SnapshotStore, register_file_tools
from agents.core.local_docs import DocumentExtraction
from agents.core.tool_rpc import ToolRPCServer

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = (
    f'xmlns:w="{W}" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
    'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


@pytest.fixture
def root(tmp_path):
    directory = tmp_path / "workspace"
    directory.mkdir()
    return directory


def body(xml):
    return ET.fromstring(f"<w:body {NS}>{xml}</w:body>")


def part(tag, xml):
    return SimpleNamespace(element=ET.fromstring(f"<w:{tag} {NS}>{xml}</w:{tag}>"))


def install_docx(monkeypatch, body_xml, rels=()):
    """A stdlib-XML ``docx`` stand-in, so the walker is tested without python-docx."""
    fake = ModuleType("docx")
    fake.__spec__ = importlib.machinery.ModuleSpec("docx", loader=None)
    seen = []

    class Document:
        def __init__(self, source):
            seen.append(source)
            self.element = SimpleNamespace(body=body(body_xml))
            self.part = SimpleNamespace(rels={
                f"rId{i}": SimpleNamespace(reltype=reltype, target_part=target)
                for i, (reltype, target) in enumerate(rels)
            })

    fake.Document = Document
    monkeypatch.setitem(sys.modules, "docx", fake)
    return seen


def p(text):
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


TABLE = (
    "<w:tbl><w:tblPr/>"
    "<w:tr><w:tc>" + p("Year") + "</w:tc><w:tc>" + p("Revenue") + "</w:tc></w:tr>"
    "<w:tr><w:tc>" + p("2025") + "</w:tc><w:tc>" + p("1,200") + "</w:tc></w:tr>"
    "<w:tr><w:tc>" + p("2026") + "</w:tc><w:tc>" + p("1,850") + "</w:tc></w:tr>"
    "</w:tbl>"
)
HEADER = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header"
FOOTER = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer"
FOOTNOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"


def test_docx_walks_tables_in_document_order_and_counts_its_parts(monkeypatch):
    textbox = (
        "<w:p><w:r><w:t>Before box</w:t></w:r><w:r><mc:AlternateContent>"
        "<mc:Choice><w:drawing><w:txbxContent>" + p("Boxed note") + "</w:txbxContent></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><w:txbxContent>" + p("Boxed note") + "</w:txbxContent></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r></w:p>"
    )
    picture = (
        "<w:p><w:r><w:drawing><a:graphic><a:graphicData><a:blip r:embed=\"rId9\"/>"
        "</a:graphicData></a:graphic></w:drawing></w:r></w:p>"
    )
    install_docx(monkeypatch, p("Intro") + TABLE + textbox + picture
                 + "<w:p><w:pPr><w:tabs><w:tab w:val=\"left\"/></w:tabs></w:pPr>"
                   "<w:r><w:t>a</w:t><w:tab/><w:t>b</w:t></w:r></w:p>"
                 + "<w:sectPr/>",
                 rels=[(HEADER, part("hdr", p("CONFIDENTIAL HEADER"))),
                       (FOOTER, part("ftr", p("Page footer"))),
                       (FOOTNOTES, part("footnotes",
                                        '<w:footnote w:type="separator" w:id="-1">' + p("---") + "</w:footnote>"
                                        '<w:footnote w:id="1">' + p("Source: audit") + "</w:footnote>"))])
    out = local_docs.extract_docx_document(io.BytesIO(b"PK"))
    assert out.text.split("\n") == [
        "[header] CONFIDENTIAL HEADER",
        "Intro",
        "| Year | Revenue |",
        "| --- | --- |",
        "| 2025 | 1,200 |",
        "| 2026 | 1,850 |",
        "Before box",
        "[text box] Boxed note",
        "",
        "a\tb",
        "[footer] Page footer",
        "[footnote 1] Source: audit",
    ]
    assert out.stopped is None and out.unit == "block"
    assert out.units_read == out.units_total == 10
    assert out.content_counts == {
        "tables": 1, "text_boxes": 1, "headers_footers": 2, "notes": 1,
        "images": 1, "charts": 0, "diagrams": 0, "embedded_objects": 0, "equations": 0,
    }


def test_docx_budget_stops_between_blocks(monkeypatch):
    install_docx(monkeypatch, p("Title") + TABLE + p("Tail"))
    cut = local_docs.extract_docx_document(io.BytesIO(b"PK"), max_text_bytes=40)
    assert cut.text == "Title\n| Year | Revenue |\n| --- | --- |" and cut.stopped == "text_limit"
    assert (cut.units_read, cut.units_total) == (2, 5)
    ticks = iter([0.0, 100.0])
    late = local_docs.extract_docx_document(io.BytesIO(b"PK"), max_seconds=50,
                                            clock=lambda: next(ticks))
    assert late.text == "Title" and late.stopped == "time_limit"


@pytest.mark.asyncio
async def test_a_table_only_docx_is_not_a_tiny_complete_read(root, monkeypatch):
    install_docx(monkeypatch, p("Intro") + TABLE,
                 rels=[(HEADER, part("hdr", p("CONFIDENTIAL HEADER")))])
    monkeypatch.setattr(file_tools, "_parser_available", lambda suffix: True)
    (root / "report.docx").write_bytes(b"PK\x03\x04 docx")
    out = await FileTools(FileScope([root])).read_file({"path": "report.docx"})
    assert out["ok"] is True and out["text_size"] > 5
    assert "| 2026 | 1,850 |" in out["content"] and "CONFIDENTIAL HEADER" in out["content"]
    assert out["document_parts"] == {"tables": 1, "headers_footers": 1}
    assert "visual_content" not in out and "coverage_warning" not in out


@pytest.mark.asyncio
async def test_docx_visual_and_unextracted_content_are_named(root, monkeypatch):
    install_docx(monkeypatch, (
        "<w:p><w:r><w:drawing><a:blip/></w:drawing></w:r>"
        "<w:r><w:pict><v:imagedata/></w:pict></w:r>"
        "<w:r><w:drawing><c:chart/></w:drawing></w:r>"
        "<w:r><mc:AlternateContent><mc:Choice><w:drawing><a:blip/></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><v:imagedata/></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
        "<w:r><w:object><v:shape><v:imagedata/></v:shape></w:object></w:r>"
        "<m:oMathPara><m:oMath><m:r><m:t>x</m:t></m:r></m:oMath></m:oMathPara></w:p>"
    ))
    monkeypatch.setattr(file_tools, "_parser_available", lambda suffix: True)
    (root / "deck.docx").write_bytes(b"PK\x03\x04 visual")
    out = await FileTools(FileScope([root])).read_file({"path": "deck.docx"})
    visual = out["visual_content"]
    assert visual["reason"] == "docx_images" and "no OCR" in visual["detail"]
    assert (visual["images"], visual["charts"], visual["diagrams"], visual["embedded_objects"]) == (3, 1, 0, 1)
    warning = out["coverage_warning"]
    assert warning["reason"] == "docx_content_not_extracted" and warning["equations"] == 1
    assert "x" not in out["content"]


def test_real_docx_with_table_header_and_picture(tmp_path):
    # python-docx is a declared runtime dependency (requirements.txt / *.lock): a
    # missing parser must fail this test, not skip it.
    import docx
    png = (b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00"
           + zlib.crc32(b"IHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00").to_bytes(4, "big")
           + len(zlib.compress(b"\x00\x00\x00\x00")).to_bytes(4, "big") + b"IDAT"
           + zlib.compress(b"\x00\x00\x00\x00")
           + zlib.crc32(b"IDAT" + zlib.compress(b"\x00\x00\x00\x00")).to_bytes(4, "big")
           + b"\x00\x00\x00\x00IEND" + zlib.crc32(b"IEND").to_bytes(4, "big"))
    (tmp_path / "dot.png").write_bytes(png)
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "CONFIDENTIAL HEADER"
    document.add_paragraph("Intro")
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([("Year", "Revenue"), ("2025", "1,200"), ("2026", "1,850")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_picture(str(tmp_path / "dot.png"))
    target = tmp_path / "real.docx"
    document.save(str(target))
    out = local_docs.extract_docx_document(io.BytesIO(target.read_bytes()))
    assert out.text.split("\n") == [
        "[header] CONFIDENTIAL HEADER", "Intro", "| Year | Revenue |", "| --- | --- |",
        "| 2025 | 1,200 |", "| 2026 | 1,850 |", "",
    ]
    assert out.content_counts["images"] == 1 and out.content_counts["tables"] == 1


REAL_NS = (
    f'xmlns:w="{W}" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def test_real_docx_footnotes_part_text_box_fallback_footer_and_tracked_changes(tmp_path):
    """Real python-docx paths the stdlib stand-in cannot exercise: a footnotes part
    python-docx has no class for (a generic blob part, parsed by ``_docx_story_root``
    through ``parse_xml``), an ``mc:AlternateContent`` text box with its VML fallback
    copy, a footer part, and tracked changes (``w:ins`` kept, ``w:delText`` dropped)."""
    import docx
    from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml.parser import parse_xml

    document = docx.Document()
    document.sections[0].footer.paragraphs[0].text = "Page footer"
    document.add_paragraph("Intro")
    sect = document.element.body[-1]  # w:sectPr stays last
    box = f'<w:txbxContent {REAL_NS}><w:p><w:r><w:t>Boxed note</w:t></w:r></w:p></w:txbxContent>'
    sect.addprevious(parse_xml(
        f'<w:p {REAL_NS}><w:r><w:t>Before box</w:t></w:r><w:r><mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData>'
        f'<wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r></w:p>'))
    sect.addprevious(parse_xml(
        f'<w:p {REAL_NS}><w:r><w:t xml:space="preserve">Net </w:t></w:r>'
        '<w:del w:id="7" w:author="A" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>gross </w:delText></w:r></w:del>'
        '<w:ins w:id="8" w:author="A" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t xml:space="preserve">adjusted </w:t></w:r></w:ins>'
        '<w:r><w:t>revenue</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r></w:p>'))
    footnotes = (
        f'<w:footnotes {REAL_NS}>'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:t>cont</w:t></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Source: audit</w:t></w:r></w:p></w:footnote>'
        '</w:footnotes>'
    ).encode()
    notes = Part(PackURI("/word/footnotes.xml"), CONTENT_TYPE.WML_FOOTNOTES, footnotes, document.part.package)
    document.part.relate_to(notes, RELATIONSHIP_TYPE.FOOTNOTES)
    target = tmp_path / "tracked.docx"
    document.save(str(target))

    reloaded = docx.Document(str(target))
    notes_part = next(rel.target_part for rel in reloaded.part.rels.values()
                      if rel.reltype == RELATIONSHIP_TYPE.FOOTNOTES)
    assert getattr(notes_part, "element", None) is None  # the blob / parse_xml branch

    out = local_docs.extract_docx_document(io.BytesIO(target.read_bytes()))
    lines = out.text.split("\n")
    assert lines == [
        "Intro",
        "Before box",
        "[text box] Boxed note",
        "Net adjusted revenue",
        "[footer] Page footer",
        "[footnote 1] Source: audit",
    ]
    assert sum(line.count("Boxed note") for line in lines) == 1
    assert "gross" not in out.text and "cont" not in out.text
    assert out.stopped is None and out.units_read == out.units_total == 5
    assert (out.content_counts["text_boxes"], out.content_counts["headers_footers"],
            out.content_counts["notes"]) == (1, 1, 1)


@pytest.mark.asyncio
async def test_a_time_budget_stop_is_not_cached_but_a_text_budget_stop_is(tmp_path, monkeypatch):
    root = tmp_path
    (root / "slow.pdf").write_bytes(b"%PDF-1.4 slow")
    monkeypatch.setattr(file_tools, "_parser_available", lambda suffix: True)
    results = iter([
        DocumentExtraction(pages=["page one"], units_total=3, units_read=1, stopped="time_limit"),
        DocumentExtraction.from_pages(["page one", "page two", "page three"]),
    ])
    calls = []

    def extract(source):
        calls.append(source)
        return next(results)

    monkeypatch.setattr(file_tools, "extract_pdf_document", extract)
    tools = FileTools(FileScope([root]))
    first = await tools.read_file({"path": "slow.pdf"})
    assert first["extraction_incomplete"]["reason"] == "time_limit" and first["cached"] is False
    assert "parses the document again" in first["extraction_incomplete"]["detail"]
    second = await tools.read_file({"path": "slow.pdf"})
    assert second["cached"] is False and "extraction_incomplete" not in second
    assert second["content"] == "page one\npage two\npage three"
    third = await tools.read_file({"path": "slow.pdf"})
    assert third["cached"] is True and len(calls) == 2

    (root / "huge.pdf").write_bytes(b"%PDF-1.4 huge")
    monkeypatch.setattr(file_tools, "extract_pdf_document", lambda source: DocumentExtraction(
        pages=["p"], units_total=9, units_read=1, stopped="text_limit"))
    assert (await tools.read_file({"path": "huge.pdf"}))["cached"] is False
    again = await tools.read_file({"path": "huge.pdf"})
    assert again["cached"] is True and again["extraction_incomplete"]["reason"] == "text_limit"


@pytest.mark.asyncio
async def test_kanban_worker_reads_share_the_document_cache(tmp_path, monkeypatch):
    from agents.core.kanban.context import KanbanContext, kanban_scope
    from agents.core.kanban.workspace_context import workspace_scope

    owner = tmp_path / "owner"
    owner.mkdir()
    tools = FileTools(FileScope([owner]), snapshots=SnapshotStore(tmp_path / "snapshots"))
    rpc = ToolRPCServer()
    register_file_tools(rpc, tools, enabled=True)
    home = tmp_path / "board-home"
    cwd = home / "kanban" / "workspaces" / "task-a"
    cwd.mkdir(parents=True)
    (cwd / "big.pdf").write_bytes(b"%PDF-1.4 worker")
    monkeypatch.setattr(file_tools, "_parser_available", lambda suffix: True)
    calls = []

    def extract(source):
        calls.append(source)
        return DocumentExtraction.from_pages(["alpha\nbeta\ngamma"])

    monkeypatch.setattr(file_tools, "extract_pdf_document", extract)
    context = KanbanContext(home, "jarvis", task_id="task-a", run_id=1, can_mutate=True)
    with kanban_scope(context), workspace_scope(cwd, still_current=lambda: True):
        first = await rpc.handle({"tool": "file_read", "args": {"path": "big.pdf", "line_limit": 1}})
        second = await rpc.handle({"tool": "file_read",
                                   "args": {"path": "big.pdf", "line_offset": 2, "line_limit": 1}})
    assert first["ok"] and first["result"]["cached"] is False
    assert second["ok"] and second["result"]["cached"] is True
    assert second["result"]["content"] == "beta\n" and len(calls) == 1


@pytest.mark.asyncio
async def test_offset_zero_beside_a_line_window_is_the_absent_offset(root):
    (root / "notes.txt").write_bytes(b"one\ntwo\nthree\n")
    tools = FileTools(FileScope([root]))
    out = await tools.read_file({"path": "notes.txt", "offset": 0, "line_offset": 2, "line_limit": 1})
    assert out["ok"] is True and out["content"] == "two\n"
    refused = await tools.read_file({"path": "notes.txt", "offset": 4, "line_offset": 2})
    assert refused == {"ok": False, "reason": "bad_window"}
    assert file_tools._line_window_args({"offset": 0, "line_offset": 2}) == (2, file_tools.DEFAULT_LINE_LIMIT)


def test_a_line_page_never_ends_inside_a_utf8_character():
    page = file_tools._line_page(io.BytesIO("é…\nnext\n".encode()), 1, 5, budget=1)
    assert page["content"] == "" and page["bytes"] == 0
    assert page["line_truncated"] is True and page["line_byte_offset"] == 0
    assert page["next_line_offset"] == 2
    whole = file_tools._line_page(io.BytesIO("é…\nnext\n".encode()), 1, 5, budget=3)
    assert whole["content"] == "é" and whole["line_truncated"] is True
