"""
local_docs.py — H12.2 Onboarding "drop folder → private chat with your docs".

One low-friction step: point Jarvis at a local folder, it indexes the documents
into memory **locally** (using whatever embedding backend is configured, with the
hash fallback so it works fully offline), and you can then chat with them. No
cloud hop, no config.

Plain-text formats (.md / .markdown / .txt) are read natively. PDF / DOCX are
best-effort: if the optional parser isn't installed the file is skipped (counted,
never fatal). The ``remember`` callable is injected so this is offline-testable.
"""

from __future__ import annotations

import asyncio
import logging
import re
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Awaitable, BinaryIO, Callable, Optional

logger = logging.getLogger("jarvis.local_docs")

TEXT_EXTS = {".md", ".markdown", ".txt", ".text", ".rst"}
DOC_EXTS = {".pdf", ".docx"}
SUPPORTED_EXTS = TEXT_EXTS | DOC_EXTS

# remember: ``async (text, metadata) -> Optional[str]``
RememberFn = Callable[[str, dict], Awaitable[Optional[str]]]


#: A parser input: a path on disk, or an in-memory binary stream (``file_read`` parses
#: the bounded bytes it read, never the path, so a file that grows after it was
#: measured cannot feed the parser more than was admitted — H444).
DocSource = Path | BinaryIO

#: H444 — the most extracted text one document read may build (UTF-8 bytes). Past it
#: extraction stops at a page boundary and says so, instead of allocating without end.
MAX_EXTRACTED_TEXT_BYTES = 32_000_000
#: H444 — the wall-clock budget for one document's extraction, checked between pages
#: (PDF) and paragraphs (DOCX). Past it extraction stops and says so.
MAX_EXTRACTION_SECONDS = 60.0
#: How many ranges a coverage report lists before it only counts the rest.
MAX_LISTED_RANGES = 16


def _parser_input(source: DocSource) -> Any:
    return str(source) if isinstance(source, Path) else source


def extract_pdf_pages(source: DocSource) -> list[str] | None:
    """Extract each PDF page's text once, or refuse the entire failed traversal."""
    try:
        import pypdf  # type: ignore

        reader = pypdf.PdfReader(_parser_input(source))
        pages: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if not isinstance(text, str):
                return None
            pages.append(text)
        return pages
    except Exception:
        return None  # Missing parser or unreadable/partially failed PDF.


@dataclass
class DocumentExtraction:
    """What one budgeted extraction produced (H444).

    ``pages`` holds the text of each page read (one entry for a DOCX). ``stopped`` is
    ``None`` when the whole document was read, else ``"text_limit"`` or
    ``"time_limit"``: the text is then a true prefix of the document, ending at a
    page (or paragraph) boundary, and ``units_read`` of ``units_total`` were read.
    ``image_pages`` are the 1-based PDF pages whose resources hold an image XObject.
    """

    pages: list[str]
    units_total: int
    units_read: int
    stopped: str | None = None
    unit: str = "page"
    image_pages: list[int] = field(default_factory=list)

    @classmethod
    def from_pages(cls, pages: list[str], *,
                   image_pages: list[int] | None = None) -> DocumentExtraction:
        return cls(pages=list(pages), units_total=len(pages), units_read=len(pages),
                   image_pages=list(image_pages or []))

    @property
    def text(self) -> str:
        return "\n".join(self.pages)


def _resolved(value: Any) -> Any:
    return value.get_object() if hasattr(value, "get_object") else value


def _has_image_xobject(page: Any) -> bool:
    """True when *page*'s resources name an image XObject (``/Subtype /Image``).

    Reads the resource dictionary only — no image is decoded. Inline images in a
    content stream and images nested inside form XObjects are not seen; a page
    whose structure cannot be read is reported as having none.
    """
    try:
        resources = _resolved(page.get("/Resources"))
        if not resources:
            return False
        xobjects = _resolved(resources.get("/XObject"))
        if not xobjects:
            return False
        return any(_resolved(ref).get("/Subtype") == "/Image" for ref in xobjects.values())
    except Exception:
        return False


def extract_pdf_document(
    source: DocSource,
    *,
    max_text_bytes: int = MAX_EXTRACTED_TEXT_BYTES,
    max_seconds: float = MAX_EXTRACTION_SECONDS,
    clock: Callable[[], float] = time.monotonic,
) -> DocumentExtraction | None:
    """A PDF's page texts within a text-size and time budget, or ``None`` on failure.

    A failed traversal still refuses as a whole (as :func:`extract_pdf_pages`). A
    budget that runs out is not a failure: the pages read so far come back with
    ``stopped`` naming the budget, never presented as the whole document. The first
    page is always attempted, so a budget can never yield an empty "success".
    """
    try:
        import pypdf  # type: ignore

        started = clock()
        reader = pypdf.PdfReader(_parser_input(source))
        page_list = reader.pages
        try:
            total: int | None = len(page_list)
        except Exception:
            total = None
        pages: list[str] = []
        image_pages: list[int] = []
        used = 0
        stopped: str | None = None
        for number, page in enumerate(page_list, 1):
            if pages and clock() - started > max_seconds:
                stopped = "time_limit"
                break
            text = page.extract_text() or ""
            if not isinstance(text, str):
                return None
            size = len(text.encode("utf-8")) + (1 if pages else 0)
            if pages and used + size > max_text_bytes:
                stopped = "text_limit"
                break
            used += size
            pages.append(text)
            if _has_image_xobject(page):
                image_pages.append(number)
        return DocumentExtraction(
            pages=pages, units_total=total if total is not None else len(pages),
            units_read=len(pages), stopped=stopped, image_pages=image_pages,
        )
    except Exception:
        return None  # Missing parser or unreadable/partially failed PDF.


def extract_docx_document(
    source: DocSource,
    *,
    max_text_bytes: int = MAX_EXTRACTED_TEXT_BYTES,
    max_seconds: float = MAX_EXTRACTION_SECONDS,
    clock: Callable[[], float] = time.monotonic,
) -> DocumentExtraction | None:
    """A DOCX's paragraph text within the same budgets, or ``None`` on failure."""
    try:
        import docx  # type: ignore

        started = clock()
        paragraphs = list(docx.Document(_parser_input(source)).paragraphs)
        kept: list[str] = []
        used = 0
        stopped: str | None = None
        for paragraph in paragraphs:
            if kept and clock() - started > max_seconds:
                stopped = "time_limit"
                break
            text = paragraph.text
            if not isinstance(text, str):
                return None
            size = len(text.encode("utf-8")) + (1 if kept else 0)
            if kept and used + size > max_text_bytes:
                stopped = "text_limit"
                break
            used += size
            kept.append(text)
        return DocumentExtraction(
            pages=["\n".join(kept)], units_total=len(paragraphs), units_read=len(kept),
            stopped=stopped, unit="paragraph",
        )
    except Exception:
        return None


#: A page needs this many non-space characters before its quality is judged, and is
#: garbled when at least this share of them is unusable (or ``(cid:N)`` glyph codes).
GARBLED_MIN_CHARS = 20
GARBLED_RATIO = 0.3
_CID = re.compile(r"\(cid:\d+\)")


def _unusable_char(char: str) -> bool:
    """The replacement mark a decoder leaves for an unmapped byte, and control,
    private-use (where a PDF font without a ToUnicode map often extracts to),
    surrogate and unassigned code points."""
    if char == "\ufffd":
        return True
    return unicodedata.category(char) in ("Cc", "Co", "Cs", "Cn")


def page_text_garbled(text: str) -> bool:
    """True when *text* is mostly characters a reader cannot use (H444).

    A heuristic over the extracted text itself: it never looks at the page image,
    so well-formed but wrong text (a bad font encoding onto real letters) is not
    caught.
    """
    cid_chars = sum(len(match) for match in _CID.findall(text))
    chars = [c for c in _CID.sub("", text) if not c.isspace()]
    considered = len(chars) + cid_chars
    if considered < GARBLED_MIN_CHARS:
        return False
    bad = sum(1 for c in chars if _unusable_char(c)) + cid_chars
    return bad / considered >= GARBLED_RATIO


def page_ranges(numbers: list[int], *, limit: int = MAX_LISTED_RANGES) -> dict:
    """Group ascending 1-based page *numbers* into ranges; list the first *limit*."""
    ranges: list[dict] = []
    count = 0
    start: int | None = None
    prev = 0
    for number in [*numbers, None]:
        if start is not None and number == prev + 1:
            prev = number
            continue
        if start is not None:
            count += 1
            if len(ranges) < limit:
                ranges.append({"start_page": start, "end_page": prev})
        if number is not None:
            start = prev = number
    omitted = count - len(ranges)
    return {"range_count": count, "ranges": ranges, "omitted_range_count": omitted,
            "ranges_truncated": omitted > 0}


def pdf_coverage_warning(pages: list[str]) -> dict | None:
    """Describe significant gaps in extracted PDF text, never inferred visuals.

    Two kinds of unusable page are counted: pages with no text at all (``gaps``,
    grouped, each naming the nearest prior text) and pages whose text is mostly
    unusable characters (``garbled_ranges``, present only when there are any; H444 —
    see :func:`page_text_garbled`).
    The warning appears when the two together pass the threshold: over 20 % of the
    pages, or ten pages or more.
    """
    total = len(pages)
    missing = 0
    gap_count = 0
    gaps: list[dict] = []
    gap_start: int | None = None
    last_text_page: int | None = None
    last_after = ""
    last_after_truncated = False
    gap_after_page: int | None = None
    gap_after = ""
    gap_after_truncated = False

    def close_gap(end_page: int) -> None:
        nonlocal gap_count
        gap_count += 1
        if len(gaps) < MAX_LISTED_RANGES:
            gaps.append({
                "start_page": gap_start,
                "end_page": end_page,
                "after_page": gap_after_page,
                "after": gap_after,
                "after_truncated": gap_after_truncated,
            })

    for number, text in enumerate(pages, 1):
        if not text.strip():
            missing += 1
            if gap_start is None:
                gap_start = number
                gap_after_page = last_text_page
                gap_after = last_after
                gap_after_truncated = last_after_truncated
            continue
        if gap_start is not None:
            close_gap(number - 1)
            gap_start = None
        trimmed = text.rstrip()
        last_text_page = number
        last_after = trimmed[-160:]
        last_after_truncated = len(trimmed) > 160
    if gap_start is not None:
        close_gap(total)
    garbled = [number for number, text in enumerate(pages, 1)
               if text.strip() and page_text_garbled(text)]
    unusable = missing + len(garbled)
    if not total or (unusable * 5 <= total and unusable < 10):
        return None
    omitted = gap_count - len(gaps)
    garbled_ranges = page_ranges(garbled)
    garbled_fields = {
        "pages_with_garbled_text": len(garbled),
        "garbled_ranges": garbled_ranges["ranges"],
        "omitted_garbled_range_count": garbled_ranges["omitted_range_count"],
    } if garbled else {}
    return {
        "reason": "pdf_text_layer_gaps",
        "detail": (
            "Some PDF pages yielded no extractable text. Do not assume the document "
            "is silent on a topic; inspect only the relevant page gaps."
        ),
        "total_pages": total,
        "pages_without_text": missing,
        "gap_count": gap_count,
        "gaps": gaps,
        "omitted_gap_count": omitted,
        "gaps_truncated": omitted > 0,
        **garbled_fields,
    }


def extract_text(path: Path) -> Optional[str]:
    """Return the document's text, or None if unreadable / parser missing."""
    ext = path.suffix.lower()
    if ext in TEXT_EXTS:
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except Exception as exc:
            logger.debug("read failed for %s: %s", path, exc)
            return None
    if ext == ".pdf":
        pages = extract_pdf_pages(path)
        return "\n".join(pages) if pages is not None else None
    if ext == ".docx":
        try:
            import docx  # type: ignore
            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        except Exception:
            return None
    return None


def chunk_text(text: str, chunk_words: int = 400, overlap: int = 40) -> list[str]:
    """Split text into ~chunk_words word windows with a small overlap."""
    words = text.split()
    if not words:
        return []
    if len(words) <= chunk_words:
        return [" ".join(words)]
    step = max(1, chunk_words - overlap)
    chunks = []
    for start in range(0, len(words), step):
        chunk = words[start:start + chunk_words]
        if chunk:
            chunks.append(" ".join(chunk))
        if start + chunk_words >= len(words):
            break
    return chunks


class LocalDocsIndexer:
    def __init__(self, remember: RememberFn, chunk_words: int = 400) -> None:
        self.remember = remember
        self.chunk_words = chunk_words

    async def index(self, folder: str | Path) -> dict:
        """Index every supported document under *folder* (recursively).

        *folder* is a **trusted** path: it comes from owner configuration, not
        from a request (the endpoint selects a pre-configured folder by key), so
        there is no request-controlled path expression here.
        """
        root = Path(folder).expanduser()
        if not root.exists() or not root.is_dir():
            return {"error": f"not a folder: {folder}"}

        # Walk + read + PDF/DOCX-parse + chunk are all blocking; run them in a
        # worker thread so indexing a large drop folder can't freeze the event
        # loop (this is awaited directly by the onboarding route). Only the async
        # remember() writes stay on the loop.
        parsed, files_skipped, skipped = await asyncio.to_thread(self._collect, root)

        files_indexed = 0
        chunks_total = 0
        for rel, chunks in parsed:
            for i, chunk in enumerate(chunks):
                await self.remember(
                    chunk,
                    {"source": "local_docs", "file": rel, "chunk": i},
                )
            files_indexed += 1
            chunks_total += len(chunks)

        return {
            "folder": str(root),
            "files_indexed": files_indexed,
            "files_skipped": files_skipped,
            "chunks": chunks_total,
            "skipped": skipped[:50],
        }

    def _collect(self, root: Path) -> tuple[list[tuple[str, list[str]]], int, list[str]]:
        """Blocking walk/extract/chunk of *root*; safe to run in a thread.

        Returns ``([(rel_path, chunks)…], files_skipped, skipped_names)``.
        """
        parsed: list[tuple[str, list[str]]] = []
        files_skipped = 0
        skipped: list[str] = []
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTS:
                continue
            text = extract_text(path)
            if not text or not text.strip():
                files_skipped += 1
                skipped.append(path.name)
                continue
            rel = str(path.relative_to(root))
            parsed.append((rel, chunk_text(text, self.chunk_words)))
        return parsed, files_skipped, skipped
